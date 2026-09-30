"""Eksport wszystkich kolumn do CSV i XLSX."""
from __future__ import annotations

import csv
from pathlib import Path

from config import EXPORT_COLUMNS


def _value(rec: dict, key: str):
    v = rec.get(key)
    if v is None:
        return ""
    if isinstance(v, bool):
        return "tak" if v else "nie"
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return ""
    return v


def export_csv(path: str | Path, records: list[dict], progress=None) -> None:
    total = max(1, len(records))
    with Path(path).open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow([h for _, h in EXPORT_COLUMNS])
        for i, rec in enumerate(records, start=1):
            w.writerow([_value(rec, k) for k, _ in EXPORT_COLUMNS])
            if progress and i % 25 == 0:
                progress(i, total)
    if progress:
        progress(total, total)


def export_xlsx(path: str | Path, records: list[dict], progress=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Wydania"

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    for c, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total = max(1, len(records))
    for r, rec in enumerate(records, start=2):
        for c, (key, _) in enumerate(EXPORT_COLUMNS, start=1):
            val = _value(rec, key)
            if key == "raw" and isinstance(val, str) and len(val) > 32000:
                val = val[:32000] + " [...]"
            ws.cell(row=r, column=c, value=val)
        if progress and (r - 1) % 25 == 0:
            progress(r - 1, total)

    widths = {"old_name": 38, "path": 55, "title": 22, "date_iso": 12, "date_raw": 26,
              "new_name": 46, "status": 20, "note": 34, "raw": 60, "model": 30}
    for c, (key, header) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(key, max(12, len(header) + 2))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(2, len(records) + 1)}"
    wb.save(str(path))
    if progress:
        progress(total, total)
